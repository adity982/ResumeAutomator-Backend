from __future__ import annotations

import io
import re
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from ..config import settings


EMAIL_PATTERN = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{8,}\d)(?!\d)")
URL_PATTERN = re.compile(r"https?://(?:www\.)?(?:linkedin\.com|github\.com)/[a-zA-Z0-9\-_/]+", re.I)


class InvalidDocument(ValueError):
    pass


def parse_pdf(content: bytes) -> dict:
    if len(content) > settings.max_upload_bytes:
        raise InvalidDocument(f"PDF exceeds the {settings.max_upload_bytes // (1024 * 1024)} MB limit.")
    if not content.startswith(b"%PDF-"):
        raise InvalidDocument("The uploaded file is not a valid PDF.")
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise InvalidDocument("The PDF is malformed or unsupported.") from exc
    try:
        if document.needs_pass:
            raise InvalidDocument("Encrypted PDFs are not supported.")
        page_count = document.page_count
        if page_count < 1 or page_count > settings.max_pdf_pages:
            raise InvalidDocument(f"PDF must contain between 1 and {settings.max_pdf_pages} pages.")
        text = "\n".join(page.get_text("text") for page in document).strip()
    finally:
        document.close()
    if len(text) < 20:
        raise InvalidDocument("The PDF contains too little selectable text. OCR is not enabled.")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    emails = EMAIL_PATTERN.findall(text)
    phones = PHONE_PATTERN.findall(text)
    links = sorted(set(URL_PATTERN.findall(text)))
    return {
        "name": lines[0][:160] if lines else "Not found",
        "email": emails[0] if emails else "Not found",
        "phone": phones[0].strip() if phones else "Not found",
        "links": links,
        "raw_content": text[:50_000],
        "page_count": page_count,
    }


THEMES = {
    "standard": ("times", 11, (0, 0, 0), (0, 0, 0)),
    "modern": ("helvetica", 11, (40, 44, 52), (0, 102, 204)),
    "executive": ("times", 11, (20, 20, 20), (0, 0, 128)),
    "minimalist": ("helvetica", 10, (70, 70, 70), (40, 44, 52)),
    "creative": ("courier", 10, (0, 0, 0), (0, 120, 110)),
}
FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"


def _unicode_fonts(pdf):
    pdf.add_font("Noto", fname=str(FONT_DIR / "NotoSans-Regular.ttf"))
    pdf.add_font("Noto", style="B", fname=str(FONT_DIR / "NotoSans-Bold.ttf"))
    pdf.add_font("Devanagari", fname=str(FONT_DIR / "NotoSansDevanagari-Regular.ttf"))
    pdf.set_fallback_fonts(["Devanagari"], exact_match=False)
    pdf.set_text_shaping(True)


def validate_glyphs(pdf, text):
    supported = set().union(*(font.cmap for font in pdf.fonts.values()))
    missing = sorted({c for c in text if not c.isspace() and ord(c) not in supported})
    if missing:
        raise InvalidDocument("The export fonts do not support these characters: " + " ".join(missing[:10]) + ". Edit them or configure an appropriate font; no characters were silently removed.")


def create_pdf(text: str, template: str, doc_type: str) -> bytes:
    font, size, text_color, header_color = THEMES.get(template, THEMES["standard"])
    cleaned = text.strip()
    if not cleaned:
        raise InvalidDocument("Document content is empty after validation.")
    pdf = FPDF()
    _unicode_fonts(pdf)
    validate_glyphs(pdf, cleaned)
    font = "Noto"
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_font(font, size=size)
    pdf.set_text_color(*text_color)
    for line in cleaned.splitlines():
        stripped = line.strip()
        if stripped and (stripped.isupper() or stripped.endswith(":")) and len(stripped) < 80:
            pdf.set_font(font, style="B", size=size + 1)
            pdf.multi_cell(0, 6, stripped, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font(font, size=size)
        else:
            pdf.multi_cell(0, 5.5, line or " ", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def create_docx(text: str, template: str, doc_type: str) -> bytes:
    document = Document()
    font, size, text_color, header_color = THEMES.get(template, THEMES["standard"])
    font_name = {"times": "Noto Serif", "helvetica": "Noto Sans", "courier": "Consolas"}[font]
    for style_name in ("Normal", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*text_color)
    for style_name in ("Heading 1", "Title"):
        document.styles[style_name].font.color.rgb = RGBColor(*header_color)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = 720000
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
        elif stripped.isupper() and len(stripped) < 80:
            document.add_heading(stripped.title(), level=1)
        elif stripped.startswith(("- ", "* ", "• ")):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(stripped)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def safe_filename(value: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(value).stem).strip("._")
    return (stem or "document")[:80]
